from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).with_name("InternMate Project Handover.docx")

BUSINESS_TABLES = [
    ("accounts_userprofile", "Role assigned to a login user.", "id, role, user_id"),
    ("operations_department", "Departments used by projects and intern profiles.", "id, created_at, updated_at, name"),
    ("operations_project", "Project master data and project schedule.", "id, created_at, updated_at, project_code, project_name, status, description, department_id, start_date, end_date"),
    ("operations_internprofile", "Intern profile, employment dates, and reporting relationships.", "id, created_at, updated_at, intern_code, internship_start_date, internship_end_date, status, department_id, manager_id, mentor_id, user_id"),
    ("operations_internprojectassignment", "Maps an intern to a project.", "id, created_at, updated_at, is_primary, assigned_from, assigned_to, intern_profile_id, project_id"),
    ("operations_mentorprojectassignment", "Maps a mentor to a project.", "id, created_at, updated_at, mentor_id, project_id"),
    ("operations_dailyprogressentry", "Core task and daily progress record.", "id, created_at, updated_at, log_date, task_name, task_description, planned_work, actual_work, completion_percentage, workflow_state, blocker_flag, blocker_description, support_needed, time_spent_hours, expected_completion_date, deliverable_link, remarks, on_time_completion_flag, created_by_id, updated_by_id, intern_profile_id, completion_status_id, final_owner_id, priority_id, task_type_id, project_id"),
    ("operations_taskattachment", "Files attached to a task update; file bytes are stored in ATP.", "id, entry_id, original_filename, content_type, content_size, content_blob, created_by_id, created_at, updated_at"),
    ("operations_aiusagerecord", "Optional AI platform and time-saved record for one task.", "id, created_at, updated_at, ai_contribution_percentage, prompt_summary, accuracy_helpfulness_rating, estimated_time_saved_hours, used_ai_flag, entry_id, ai_assistance_type_id, ai_tool_id, human_review_status_id, output_usage_id"),
    ("operations_mentorreviewcomment", "Mentor review and workflow comments on a task.", "id, created_at, updated_at, action, comment_text, action_at, entry_id, reviewer_id"),
    ("operations_taskdependency", "Dependencies and their status for a task.", "id, created_at, updated_at, dependency_text, dependency_status, entry_id"),
    ("operations_auditlog", "Change history for audited entities.", "id, created_at, updated_at, entity_type, entity_id, action_type, old_value_json, new_value_json, actor_id"),
    ("operations_masterdatalist", "Master-data list definition, such as Priority or AI Tool.", "id, created_at, updated_at, name, description"),
    ("operations_masterdataitem", "Individual selectable master-data values.", "id, created_at, updated_at, value, display_order, is_active, list_ref_id"),
    ("operations_skilldomain", "Skill-domain master records.", "id, created_at, updated_at, name"),
    ("operations_internprofile_skills", "Maps an intern profile to a skill domain.", "id, internprofile_id, skilldomain_id"),
]

DJANGO_TABLES = [
    ("auth_user", "User login, password hash, names, email, and active/staff flags.", "id, password, last_login, is_superuser, username, last_name, email, is_staff, is_active, date_joined, first_name"),
    ("auth_group", "Django role/permission group.", "id, name"),
    ("auth_permission", "Django permission definition.", "id, name, content_type_id, codename"),
    ("auth_group_permissions", "Maps a group to permissions.", "id, group_id, permission_id"),
    ("auth_user_groups", "Maps a user to groups.", "id, user_id, group_id"),
    ("auth_user_user_permissions", "Maps a user to direct permissions.", "id, user_id, permission_id"),
]


def shade(cell, color):
    props = cell._tc.get_or_add_tcPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:fill"), color)
    props.append(fill)


def set_cell_text(cell, value, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(value)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(size)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, rows):
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.65), Inches(1.9), Inches(3.65)]
    headers = ["Table", "Purpose", "Columns"]
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = widths[index]
        shade(cell, "17365D")
        set_cell_text(cell, text, bold=True, size=9)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(table.rows[0])
    for number, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = widths[index]
            if number % 2:
                shade(cells[index], "F2F6FA")
            set_cell_text(cells[index], value, bold=index == 0, size=8.5)
    doc.add_paragraph()


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.name = "Aptos Display"
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    for run in p.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(10.5)


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("InternMate Project Handover")
    run.font.name = "Aptos Display"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    subtitle = doc.add_paragraph("Application overview, deployment guide, roles, and Oracle ATP data model")
    subtitle.paragraph_format.space_after = Pt(16)
    for run in subtitle.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(70, 70, 70)

    add_heading(doc, "Executive Summary")
    add_body(doc, "InternMate is a role-based work-management application for assigning intern tasks, recording progress, tracking blockers, capturing optional AI assistance and time saved, and producing team performance views. The current production-style configuration stores application data in Oracle Autonomous Transaction Processing (ATP).")
    add_body(doc, "The active web application is React served by a FastAPI backend. It reads and writes the existing InternMate Oracle ATP schema. The repository still contains Django apps and migration files because the data model and authentication tables originated in Django; they remain useful for maintenance and compatibility, but the current running web server is FastAPI, not Django runserver.")

    add_heading(doc, "Technology and Runtime")
    add_bullet(doc, "Frontend: React with Vite and Bootstrap styling. The compiled frontend is served by FastAPI.")
    add_bullet(doc, "Backend: Python FastAPI with Uvicorn, SQLAlchemy, and python-oracledb.")
    add_bullet(doc, "Database: Oracle Autonomous Transaction Processing. The live application uses the Oracle connection when INTERNMATE_DATABASE is set to oracle.")
    add_bullet(doc, "Exports: CSV, XLSX through openpyxl, and PDF through ReportLab.")
    add_bullet(doc, "Legacy source: Django apps, migrations, templates, and auth-table naming are retained for data-model history and compatibility.")

    add_heading(doc, "What the Django Files Are")
    add_body(doc, "The folders accounts, operations, reporting, internmate, templates, and manage.py are the original Django project structure. They define the original models, migrations, permissions, and server-rendered pages. Do not delete them without a planned migration review because the ATP schema still uses Django-style table names such as auth_user and operations_dailyprogressentry.")
    add_body(doc, "For the deployed application, start FastAPI with backend.app.main:app. The Django files are not the active HTTP server, but they remain the reference for the original domain model and can be useful when tracing database fields.")

    add_heading(doc, "Roles and Main Workflows")
    role_table = doc.add_table(rows=1, cols=3)
    role_table.style = "Table Grid"
    for i, label in enumerate(["Role", "Main actions", "Views and restrictions"]):
        shade(role_table.rows[0].cells[i], "17365D")
        set_cell_text(role_table.rows[0].cells[i], label, bold=True, size=9)
        for run in role_table.rows[0].cells[i].paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
    set_repeat_table_header(role_table.rows[0])
    for num, row in enumerate([
        ("Admin", "Create projects; manage project and master-data context; review all work.", "Dashboard, Daily Logs, Daily Stats, Reports, CODEX CLI Guide, and Admin page."),
        ("Mentor", "Assign tasks to an intern and project; monitor intern work updates.", "Dashboard, Daily Logs, Daily Stats, and Reports. Sees workspace task updates."),
        ("Intern", "Select an intern/project task, submit work completed, completion, hours, blockers, attachment, AI platform, and time saved.", "Dashboard, My Tasks, Daily Stats, and Reports. Current shared workspace mode exposes the configured intern/task lists."),
    ]):
        cells = role_table.add_row().cells
        for i, value in enumerate(row):
            if num % 2:
                shade(cells[i], "F2F6FA")
            set_cell_text(cells[i], value, bold=i == 0, size=9)
    doc.add_paragraph()

    add_heading(doc, "How the Pages Differ")
    add_bullet(doc, "Dashboard: KPI overview and the latest task activity. It aggregates completion, blockers, AI-assisted task percentage, and time saved.")
    add_bullet(doc, "Daily Logs or My Tasks: one row per task record. This is the operational list used to inspect and update individual work.")
    add_bullet(doc, "Daily Stats: groups records by date to show daily workload, average completion, AI usage, hours spent, and time saved.")
    add_bullet(doc, "Reports: groups records by intern to compare workload, completion, blocked work, and time saved; supports CSV, Excel, and PDF export.")

    add_heading(doc, "Oracle ATP Database Dependency")
    add_body(doc, "Yes, the current local running application is connected to ATP. User records, projects, assignments, task updates, attachments, AI usage, and dashboard totals are read from and written to ATP. A task update is saved first; the application reloads the same ATP-backed daily-log data, and the Dashboard, Daily Stats, and Reports calculate their values from those rows.")
    add_body(doc, "The code also includes a SQLite fallback for development if the environment variable INTERNMATE_DATABASE is set to sqlite. For deployment against ATP, keep INTERNMATE_DATABASE=oracle and provide the Oracle username, password, DSN, and wallet directory as protected environment variables. Do not place credentials or wallet files inside a shared ZIP, Git repository, or public file share.")

    add_heading(doc, "Business Data Tables")
    add_body(doc, "The following table inventory was read from the current Oracle ATP schema. Foreign-key columns end in _id and connect the records shown in the purpose column.")
    add_table(doc, BUSINESS_TABLES)

    add_heading(doc, "Authentication and Django Framework Tables")
    add_body(doc, "These tables support login, roles, groups, and permission compatibility. Django also maintains framework tables such as django_migrations and django_content_type; they are standard framework metadata and are not day-to-day business data.")
    add_table(doc, DJANGO_TABLES)

    add_heading(doc, "Deployment Handover")
    add_bullet(doc, "Copy the project source without virtual environments, node_modules, .env files, Oracle wallets, or private keys. Install dependencies on the target server.")
    add_bullet(doc, "Build the frontend in frontend with npm run build. FastAPI serves the generated frontend/dist files.")
    add_bullet(doc, "Create a protected .env on the server with the ATP connection values and INTERNMATE_DATABASE=oracle.")
    add_bullet(doc, "Run the application with: .venv-win\\Scripts\\python.exe -m uvicorn backend.app.main:app --host 0.0.0.0 --port 8001 on Windows, or the equivalent virtual-environment Python command on Linux.")
    add_bullet(doc, "For a public deployment, place a reverse proxy and HTTPS certificate in front of Uvicorn, restrict ATP wallet access, and do not expose the database directly to the internet.")

    add_heading(doc, "Key Repository Locations")
    add_bullet(doc, "backend/app/main.py: active FastAPI routes, ATP database connection, authentication, and API logic.")
    add_bullet(doc, "frontend/src/roleClassicApp.jsx: current role-aware React application UI.")
    add_bullet(doc, "frontend/dist: compiled frontend served by the backend after npm run build.")
    add_bullet(doc, "operations/models.py and operations/migrations: original Django domain model and database history.")
    add_bullet(doc, ".env.example: example configuration structure only. Replace values on the deployment server; never share live secrets.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("InternMate project handover")
    for run in footer.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(100, 100, 100)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
